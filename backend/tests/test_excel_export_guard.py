"""P2：①公式注入消毒 ②零可识别 JSON 时明确报错而非 IndexError"""
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

import pytest


def _mk_core_events(d: Path, event: str):
    d.mkdir(parents=True, exist_ok=True)
    (d / "core_events_aggregated.json").write_text(json.dumps(
        {"chapter_1": [{"id": "e1", "event": event, "characters": "张三", "function": "铺垫"}]},
        ensure_ascii=False), encoding="utf-8")


def test_formula_like_event_is_sanitized(tmp_path):
    from openpyxl import load_workbook
    from backend.utils.excel_export import export_aggregated_to_excel as export_to_excel
    d = tmp_path / "agg"
    _mk_core_events(d, "=cmd|'/c calc'!A1")
    out = export_to_excel(d, tmp_path / "a.xlsx")

    wb = load_workbook(out)
    ws = wb["核心事件"]
    cell = ws.cell(row=2, column=3)
    assert cell.data_type != "f", "公式串必须被转义为文本"
    assert str(cell.value).startswith("'")


def test_no_recognized_json_raises_clear_error(tmp_path):
    from backend.utils.excel_export import export_aggregated_to_excel as export_to_excel
    d = tmp_path / "agg"
    d.mkdir()
    (d / "unrecognized.json").write_text("{}", encoding="utf-8")
    with pytest.raises(ValueError, match="可识别"):
        export_to_excel(d, tmp_path / "b.xlsx")
